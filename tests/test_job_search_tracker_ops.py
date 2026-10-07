from __future__ import annotations

import importlib.util
import io
import json
import sys
import tempfile
import unittest
from contextlib import redirect_stdout
from pathlib import Path

try:
    import openpyxl  # noqa: F401
except ImportError:  # pragma: no cover
    openpyxl = None

SCRIPT = Path(__file__).resolve().parents[1] / "skills" / "job-search-pipeline" / "scripts" / "tracker_ops.py"


def load():
    spec = importlib.util.spec_from_file_location("tracker_ops", SCRIPT)
    mod = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = mod
    spec.loader.exec_module(mod)
    return mod


@unittest.skipIf(openpyxl is None, "openpyxl not installed")
class TrackerOpsTests(unittest.TestCase):
    def setUp(self) -> None:
        self.ops = load()
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        self.tracker = str(self.dir / "tracker.xlsx")
        self.assertEqual(self.run_ops("init", "--tracker", self.tracker)[0], 0)
        jobs = [
            {"company": "Acme", "title": "Staff Engineer", "link": "https://www.linkedin.com/jobs/view/111/",
             "fit": 80, "tier": "A", "required": ["Java", "Kafka", "GraphQL"], "nice": ["Rust"]},
            {"company": "Globex", "title": "Senior Backend", "link": "https://boards.example.com/jobs?gh_jid=222",
             "fit": 70, "tier": "B", "required": ["Java", "GraphQL"], "red_flags": ["pay not stated"]},
        ]
        self.jobs = self.dir / "jobs.json"
        self.jobs.write_text(json.dumps({"jobs": jobs}), encoding="utf-8")

    def tearDown(self) -> None:
        self.tmp.cleanup()

    def run_ops(self, *args):
        buf = io.StringIO()
        with redirect_stdout(buf):
            code = self.ops.main(list(args))
        return code, (json.loads(buf.getvalue()) if buf.getvalue().strip() else None)

    def test_append_dedups_and_lists_known(self) -> None:
        code, out = self.run_ops("append", "--tracker", self.tracker, "--jobs", str(self.jobs), "--cv-skill", "^java$")
        self.assertEqual(code, 0)
        self.assertEqual(len(out["added"]), 2)
        code, out = self.run_ops("append", "--tracker", self.tracker, "--jobs", str(self.jobs))
        self.assertEqual(out["added"], [])
        self.assertEqual(len(out["skipped_duplicates"]), 2)
        _, known = self.run_ops("known", "--tracker", self.tracker)
        self.assertIn("111", known["ids"])
        self.assertIn("222", known["ids"])

    def test_expect_mtime_mismatch_aborts_without_writing(self) -> None:
        code, _ = self.run_ops("append", "--tracker", self.tracker, "--jobs", str(self.jobs), "--expect-mtime", "1")
        self.assertEqual(code, 3)
        _, known = self.run_ops("known", "--tracker", self.tracker)
        self.assertEqual(known["ids"], [])

    def test_approvals_and_moves(self) -> None:
        self.run_ops("append", "--tracker", self.tracker, "--jobs", str(self.jobs))
        wb = openpyxl.load_workbook(self.tracker)
        wb["For Approval"]["B2"] = "☑ Approve"
        wb["For Approval"]["B3"] = "✖ Reject"
        wb.save(self.tracker)
        _, appr = self.run_ops("approvals", "--tracker", self.tracker)
        self.assertEqual([x["company"] for x in appr["approve"]], ["Acme"])
        self.assertEqual([x["company"] for x in appr["reject"]], ["Globex"])
        link = "https://www.linkedin.com/jobs/view/111/"
        self.assertEqual(self.run_ops("move", "--tracker", self.tracker, "--link", link, "--to", "Applied",
                                      "--date", "2026-01-05")[0], 0)
        self.assertEqual(self.run_ops("move", "--tracker", self.tracker, "--link", link, "--to", "Closed",
                                      "--date", "2026-01-20", "--outcome", "Rejected", "--detail", "not moving forward")[0], 0)
        self.assertEqual(self.run_ops("move", "--tracker", self.tracker, "--link",
                                      "https://boards.example.com/jobs?gh_jid=222", "--to", "Excluded")[0], 0)
        wb = openpyxl.load_workbook(self.tracker)
        closed = [c.value for c in wb["Closed"][2]]
        self.assertIn("Rejected", closed)
        self.assertIn("2026-01-05 Applied; 2026-01-20 Closed (not moving forward)", closed)
        self.assertEqual(wb["Excluded"]["A2"].value, "Globex")
        self.assertIsNone(wb["For Approval"]["F2"].value)

    def test_write_returns_mtime_for_chained_writes_and_refuses_backward_moves(self) -> None:
        _, out = self.run_ops("append", "--tracker", self.tracker, "--jobs", str(self.jobs))
        link = "https://www.linkedin.com/jobs/view/111/"
        code, out = self.run_ops("move", "--tracker", self.tracker, "--link", link, "--to", "Interviewing",
                                 "--expect-mtime", str(out["mtime"]), "--interview-stage", "Screening")
        self.assertEqual(code, 0)
        code, _ = self.run_ops("move", "--tracker", self.tracker, "--link", link, "--to", "Applied",
                               "--expect-mtime", str(out["mtime"]))
        self.assertEqual(code, 2)
        excluded = self.dir / "excluded.json"
        excluded.write_text(json.dumps([{"company": "Initech", "title": "Dev", "link": "https://x.example/j/9",
                                         "reason": "On-site only", "typo_key": 1}]), encoding="utf-8")
        code, out = self.run_ops("exclude", "--tracker", self.tracker, "--jobs", str(excluded),
                                 "--expect-mtime", str(out["mtime"]))
        self.assertEqual((code, out["excluded"]), (0, ["https://x.example/j/9"]))

    def test_sidegig_upserts_by_platform(self) -> None:
        data = self.dir / "gigs.json"
        data.write_text(json.dumps([{"platform": "ExampleGigs", "status": "Invited", "next_action": "Take test"}]),
                        encoding="utf-8")
        _, out = self.run_ops("sidegig", "--tracker", self.tracker, "--data", str(data))
        self.assertEqual(out["added"], ["ExampleGigs"])
        data.write_text(json.dumps([{"platform": "examplegigs", "status": "Active"}]), encoding="utf-8")
        _, out = self.run_ops("sidegig", "--tracker", self.tracker, "--data", str(data))
        self.assertEqual(out["updated"], ["examplegigs"])
        wb = openpyxl.load_workbook(self.tracker)
        self.assertEqual(wb["Side Gigs"]["D2"].value, "Active")
        self.assertIsNone(wb["Side Gigs"]["B3"].value)

    def test_gaps_flags_repeated_gap_and_keeps_owner_status(self) -> None:
        self.run_ops("append", "--tracker", self.tracker, "--jobs", str(self.jobs))
        rules = self.dir / "rules.json"
        rules.write_text(json.dumps({"cv_skills": ["^java$"], "rules": [{"match": "kafka", "ref": "P-1", "week": "W2"}]}),
                         encoding="utf-8")
        wb = openpyxl.load_workbook(self.tracker)
        wb["Skill Map"]["F5"] = "Owner note"  # Rust row: owner-set status must survive
        wb.save(self.tracker)
        code, out = self.run_ops("gaps", "--tracker", self.tracker, "--rules", str(rules))
        self.assertEqual(code, 0)
        self.assertEqual(out["repeated_gaps"], [{"skill": "GraphQL", "jobs": 2}])
        wb = openpyxl.load_workbook(self.tracker)
        self.assertEqual(wb["Skill Map"]["F5"].value, "Owner note")
        fa = wb["For Approval"]
        hdr = [c.value for c in fa[1]]
        self.assertEqual(fa.cell(2, hdr.index("Readiness %") + 1).value, 0.38)
        self.assertTrue(str(fa.cell(2, hdr.index("Next study action") + 1).value).startswith("Close GAP: GraphQL"))


if __name__ == "__main__":
    unittest.main()
