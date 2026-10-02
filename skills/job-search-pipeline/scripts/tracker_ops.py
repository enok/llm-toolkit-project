#!/usr/bin/env python3
"""Generic job-search tracker operations (xlsx, openpyxl).

Subcommands: init, known, approvals, append, exclude, move, gaps, sidegig.
Write subcommands print the new "mtime"; pass it as --expect-mtime to the next
write. Limitation: openpyxl does not keep charts, images, pivots, or in-cell
checkbox extensions, and delete_rows does not shift hyperlinks or conditional
formats - keep the tracker to plain cells, formats, and list validations.
Every write is guarded: the workbook's mtime is captured on load and
re-checked right before an atomic replace, so a concurrent edit by the
owner aborts the write (exit 3) instead of being overwritten. Sheets the
tool does not know about are preserved untouched.

Exit codes: 0 ok, 2 usage/config error, 3 file changed during the run
(or --expect-mtime mismatch), 4 file locked / not writable.
Runs on Windows, Linux and macOS with Python 3.9+ and openpyxl.
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import os
import re
import shutil
import sys
from copy import copy
from pathlib import Path

try:
    import openpyxl
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:  # pragma: no cover - environment guard
    sys.stderr.write("openpyxl is required: python -m pip install openpyxl\n")
    sys.exit(2)

APPROVE, REJECT, LATER, PENDING = "☑ Approve", "✖ Reject", "⏸ Later", "☐"
MARKS = [PENDING, APPROVE, REJECT, LATER]
PIPELINE = ["For Approval", "Applied", "Interviewing", "Closed"]

SCHEMA = {
    "Summary": ["Job search tracker", None],
    "For Approval": ["#", "Approve ☐", "Decision notes", "Tier", "Fit score", "Company", "Title", "Job link",
                     "Posted", "Originally posted", "Location / eligibility", "Employment type", "Pay (as posted)",
                     "Currency", "Est. USD/mo min", "Est. USD/mo max", "Meets pay rule", "Apply method",
                     "Years required", "English", "Required skills", "Nice to have", "Must-have gaps",
                     "Readiness %", "Covered", "In study plan", "Missing (not in plan)", "Next study action",
                     "Red flags", "Summary"],
    "Applied": ["#", "Company", "Title", "Job link", "Channel", "Date applied", "Latest status", "Pipeline stage",
                "Stage history", "Next action", "Follow-up date", "Status date", "Pay", "Currency",
                "Est USD/month", "Fit", "Required skills", "Readiness %", "Covered", "In study plan", "Missing",
                "Next interview-prep action", "Mail link", "Notes"],
    "Interviewing": ["#", "Company", "Title", "Job link", "Channel", "Date applied", "Latest status",
                     "Pipeline stage", "Stage history", "Next action", "Follow-up date", "Status date", "Pay",
                     "Currency", "Est USD/month", "Fit", "Required skills", "Readiness %", "Covered",
                     "In study plan", "Missing", "Next interview-prep action", "Mail link", "Notes",
                     "Interview stage", "Next interview date", "Interviewers", "Prep material"],
    "Closed": ["#", "Company", "Title", "Job link", "Date applied", "Closed date", "Outcome", "Outcome detail",
               "Stage reached", "Stage history", "Lessons / notes"],
    "Side Gigs": ["#", "Platform", "Link", "My status", "Roles/projects applied", "Pay (USD/hr)",
                  "Meets my floor", "Eligible from my country", "Payment method", "Next action",
                  "Next action date", "Last activity", "Mail link", "Notes/sources"],
    "Excluded": ["Company", "Title", "Job link", "Reason", "Fit", "Pay"],
    "Skill Map": ["Company", "Title", "Bucket", "Skill", "Required/Nice", "Status", "Study plan ref", "Week/phase"],
    "Gap Summary": ["Skill", "# jobs requiring", "Status", "Study plan ref", "Suggested action"],
}
# Header/sheet aliases so trackers created by earlier versions keep working.
ALIASES = {"Mail link": ["Gmail link"], "Side Gigs": ["AI Gigs"], "Approve ☐": ["Approve", "Approve?"],
           "Missing (not in plan)": ["Missing"], "Est. USD/mo max": ["Est USD/month"], "Fit score": ["Fit"],
           "Pay (as posted)": ["Pay"]}
DEFAULT_ID_PATTERNS = [r"/jobs/view/(\d+)", r"[?&](?:currentJobId|jk|gh_jid|jobId)=([\w-]+)"]
CV = "Experience (CV)"
MTIME_TOLERANCE = 0.01
STAGE_ORDER = {"For Approval": 0, "Applied": 1, "Interviewing": 2, "Closed": 3}


class Guarded(Exception):
    def __init__(self, code: int, msg: str):
        super().__init__(msg)
        self.code = code


def today() -> str:
    return dt.date.today().isoformat()


def sheet(wb, name):
    for n in [name] + ALIASES.get(name, []):
        if n in wb.sheetnames:
            return wb[n]
    return None


def headers(ws) -> dict:
    out = {}
    for c in ws[1]:
        if c.value not in (None, ""):
            out[str(c.value).strip()] = c.column
    for canon, alts in ALIASES.items():
        for a in alts:
            if a in out and canon not in out:
                out[canon] = out[a]
    return out


def last_row(ws, col: int) -> int:
    r = ws.max_row
    while r > 1 and ws.cell(r, col).value in (None, ""):
        r -= 1
    return r


def copy_row_style(ws, src: int, dst: int) -> None:
    if src < 2:
        return
    for c in range(1, ws.max_column + 1):
        s, d = ws.cell(src, c), ws.cell(dst, c)
        if s.has_style:
            d.font, d.border, d.fill = copy(s.font), copy(s.border), copy(s.fill)
            d.alignment, d.number_format = copy(s.alignment), s.number_format


def put(ws, h: dict, row: int, values: dict, clear: bool = False) -> None:
    for k, v in values.items():
        if k in h and (v is not None or clear):
            ws.cell(row, h[k]).value = v


def get(ws, h: dict, row: int, name: str):
    return ws.cell(row, h[name]).value if name in h else None


def job_key(link, patterns) -> str | None:
    s = str(link or "").strip()
    if not s:
        return None
    for p in patterns:
        m = re.search(p, s)
        if m:
            return m.group(1)
    return re.sub(r"[?#].*$", "", s).rstrip("/").lower()


def ct_key(company, title) -> str:
    def norm(x):
        return re.sub(r"\s+", " ", str(x or "")).strip().lower()
    return f"{norm(company)} | {norm(title)}"


class Tracker:
    def __init__(self, path: str, expect_mtime: float | None = None, backup_dir: str | None = None):
        self.path = Path(path)
        if not self.path.exists():
            raise Guarded(2, f"tracker not found: {self.path}")
        self.mtime = self.path.stat().st_mtime
        if expect_mtime is not None and abs(self.mtime - expect_mtime) > MTIME_TOLERANCE:
            raise Guarded(3, f"tracker mtime {self.mtime} != expected {expect_mtime}; it was edited (or re-synced "
                             "by a cloud drive) since the last read - re-read first")
        self.backup_dir = backup_dir
        self.wb = openpyxl.load_workbook(self.path)
        self.sheets_before = list(self.wb.sheetnames)

    def save(self, dry_run: bool = False) -> None:
        if dry_run:
            return
        if list(self.wb.sheetnames)[: len(self.sheets_before)] != self.sheets_before:
            raise Guarded(2, "refusing to save: an existing sheet was removed or reordered")
        if abs(self.path.stat().st_mtime - self.mtime) > MTIME_TOLERANCE:
            raise Guarded(3, "tracker changed on disk during this run; nothing written")
        if self.backup_dir:
            bd = Path(self.backup_dir)
            bd.mkdir(parents=True, exist_ok=True)
            stamp = dt.datetime.now().strftime("%Y%m%d-%H%M%S")
            shutil.copy2(self.path, bd / f"{self.path.stem}.bak-{stamp}{self.path.suffix}")
        tmp = self.path.with_name(self.path.name + ".tmp-tracker-ops")
        try:
            self.wb.save(tmp)
            os.replace(tmp, self.path)
        except PermissionError as e:
            raise Guarded(4, f"tracker is locked or read-only (close it in the spreadsheet app?): {e}")
        finally:
            if tmp.exists():
                tmp.unlink()
        self.mtime = self.path.stat().st_mtime

    def result(self, payload: dict, dry_run: bool) -> dict:
        payload.update({"dry_run": dry_run, "mtime": self.mtime})
        return payload

    def ws(self, name: str):
        w = sheet(self.wb, name)
        if w is None:
            raise Guarded(2, f"sheet '{name}' not found; run 'init' on a new tracker or add the sheet")
        return w

    def known(self, patterns) -> tuple[set, set]:
        ids, cts = set(), set()
        for name in PIPELINE + ["Excluded"]:
            ws = sheet(self.wb, name)
            if ws is None:
                continue
            h = headers(ws)
            if "Company" not in h:
                continue
            for r in range(2, ws.max_row + 1):
                co = get(ws, h, r, "Company")
                if co in (None, ""):
                    continue
                k = job_key(get(ws, h, r, "Job link"), patterns)
                if k:
                    ids.add(k)
                cts.add(ct_key(co, get(ws, h, r, "Title")))
        return ids, cts

    def find(self, link: str, patterns, names=PIPELINE):
        want = job_key(link, patterns)
        for name in names:
            ws = sheet(self.wb, name)
            if ws is None:
                continue
            h = headers(ws)
            for r in range(2, ws.max_row + 1):
                if job_key(get(ws, h, r, "Job link"), patterns) == want:
                    return ws, r
        return None, None


def cmd_init(a) -> dict:
    p = Path(a.tracker)
    if p.exists() and not a.force:
        raise Guarded(2, f"{p} exists; pass --force to overwrite")
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, hdr in SCHEMA.items():
        ws = wb.create_sheet(name)
        ws.append(hdr)
        ws.freeze_panes = "A2"
        for c in ws[1]:
            c.font = openpyxl.styles.Font(bold=True)
    s = wb["Summary"]
    for line in ["Pipeline: For Approval -> Applied -> Interviewing -> Closed (+ Excluded, Side Gigs).",
                 f"Approve column marks: {', '.join(MARKS)}. Only marks the owner sets count as approval.",
                 "Skill Map Status: Experience (CV) | Study plan ... | GAP - not in plan. Gap Summary is rebuilt by 'gaps'."]:
        s.append([line])
    dv = DataValidation(type="list", formula1='"' + ",".join(MARKS) + '"', allow_blank=True)
    wb["For Approval"].add_data_validation(dv)
    dv.add("B2:B1000")
    p.parent.mkdir(parents=True, exist_ok=True)
    wb.save(p)
    return {"created": str(p), "sheets": wb.sheetnames}


def cmd_known(a) -> dict:
    t = Tracker(a.tracker)
    ids, cts = t.known(a.id_pattern or DEFAULT_ID_PATTERNS)
    out = {"mtime": t.mtime, "ids": sorted(ids), "company_title": sorted(cts)}
    if a.out:
        Path(a.out).parent.mkdir(parents=True, exist_ok=True)
        Path(a.out).write_text(json.dumps(out, ensure_ascii=False, indent=1), encoding="utf-8")
        return {"mtime": t.mtime, "ids": len(ids), "company_title": len(cts), "out": a.out}
    return out


def cmd_approvals(a) -> dict:
    t = Tracker(a.tracker)
    ws = t.ws("For Approval")
    h = headers(ws)
    col = h.get("Approve ☐", 2)
    res = {"approve": [], "reject": [], "later": [], "pending": [], "mtime": t.mtime}
    for r in range(2, ws.max_row + 1):
        if get(ws, h, r, "Company") in (None, ""):
            continue
        v = ws.cell(r, col).value
        sv = str(v).strip() if v is not None else ""
        bucket = ("approve" if v is True or sv.startswith("☑") or sv.upper() == "TRUE" else
                  "reject" if sv.startswith("✖") else "later" if sv.startswith("⏸") else "pending")
        res[bucket].append({"row": r, "company": get(ws, h, r, "Company"), "title": get(ws, h, r, "Title"),
                            "link": get(ws, h, r, "Job link"), "notes": get(ws, h, r, "Decision notes"),
                            "apply_method": get(ws, h, r, "Apply method")})
    return res


def load_jobs(path: str) -> list:
    data = json.loads(Path(path).read_text(encoding="utf-8"))
    jobs = data["jobs"] if isinstance(data, dict) else data
    for j in jobs:
        for req in ("company", "title", "link"):
            if not j.get(req):
                raise Guarded(2, f"job entry missing '{req}' (required: company, title, link): {j}")
    return jobs


def ignored_keys(jobs: list, known: set) -> list:
    return sorted({k for j in jobs for k in j if k not in known})


FIELD_MAP = {"tier": "Tier", "fit": "Fit score", "company": "Company", "title": "Title", "link": "Job link",
             "posted": "Posted", "originally_posted": "Originally posted", "location": "Location / eligibility",
             "employment_type": "Employment type", "pay": "Pay (as posted)", "currency": "Currency",
             "usd_month_min": "Est. USD/mo min", "usd_month_max": "Est. USD/mo max", "meets_pay": "Meets pay rule",
             "apply_method": "Apply method", "years": "Years required", "english": "English",
             "red_flags": "Red flags", "summary": "Summary", "notes": "Decision notes"}


def cmd_append(a) -> dict:
    t = Tracker(a.tracker, a.expect_mtime, a.backup_dir)
    pats = a.id_pattern or DEFAULT_ID_PATTERNS
    ids, cts = t.known(pats)
    fa, sm = t.ws("For Approval"), t.ws("Skill Map")
    hf, hs = headers(fa), headers(sm)
    r = last_row(fa, hf.get("Company", 6))
    num = get(fa, hf, r, "#") if r > 1 else 0
    num = num if isinstance(num, int) else r - 1
    sr = last_row(sm, hs.get("Company", 1))
    added, skipped = [], []
    jobs = load_jobs(a.jobs)
    cv_res = [re.compile(x, re.I) for x in (a.cv_skill or [])]
    for j in jobs:
        k, c = job_key(j["link"], pats), ct_key(j["company"], j["title"])
        if k in ids or c in cts:
            skipped.append(j["link"])
            continue
        r += 1
        num += 1
        copy_row_style(fa, r - 1, r)
        vals = {FIELD_MAP[f]: v for f, v in j.items() if f in FIELD_MAP}
        if isinstance(vals.get("Red flags"), list):
            vals["Red flags"] = "; ".join(vals["Red flags"])
        vals.update({"#": num, "Approve ☐": PENDING, "Decision notes": j.get("notes") or a.note,
                     "Required skills": "; ".join(j.get("required", [])) or None,
                     "Nice to have": "; ".join(j.get("nice", [])) or None})
        put(fa, hf, r, vals)
        for kind, lst in (("Required", j.get("required", [])), ("Nice", j.get("nice", []))):
            for skill in lst:
                sr += 1
                copy_row_style(sm, sr - 1, sr)
                status = CV if any(rx.search(skill) for rx in cv_res) else None
                put(sm, hs, sr, {"Company": j["company"], "Title": j["title"], "Bucket": "For Approval",
                                 "Skill": skill, "Required/Nice": kind, "Status": status})
        ids.add(k)
        cts.add(c)
        added.append({"row": r, "company": j["company"], "title": j["title"]})
    t.save(a.dry_run)
    return t.result({"added": added, "skipped_duplicates": skipped,
                     "ignored_keys": ignored_keys(jobs, set(FIELD_MAP) | {"required", "nice"})}, a.dry_run)


def cmd_exclude(a) -> dict:
    t = Tracker(a.tracker, a.expect_mtime, a.backup_dir)
    pats = a.id_pattern or DEFAULT_ID_PATTERNS
    ids, cts = t.known(pats)
    ex = t.ws("Excluded")
    h = headers(ex)
    r = last_row(ex, h.get("Company", 1))
    added = []
    for j in load_jobs(a.jobs):
        if job_key(j["link"], pats) in ids or ct_key(j["company"], j["title"]) in cts:
            continue
        r += 1
        copy_row_style(ex, r - 1, r)
        put(ex, h, r, {"Company": j["company"], "Title": j["title"], "Job link": j["link"],
                       "Reason": j.get("reason", "Excluded"), "Fit": j.get("fit"), "Pay": j.get("pay")})
        added.append(j["link"])
    t.save(a.dry_run)
    return t.result({"excluded": added}, a.dry_run)


def add_days(iso: str, days: int) -> str:
    return (dt.date.fromisoformat(iso) + dt.timedelta(days=days)).isoformat()


def cmd_move(a) -> dict:
    """Move one job row along the pipeline (or to Excluded), carrying shared columns by header name."""
    t = Tracker(a.tracker, a.expect_mtime, a.backup_dir)
    pats = a.id_pattern or DEFAULT_ID_PATTERNS
    src, r = t.find(a.link, pats)
    if src is None:
        raise Guarded(2, f"job not found in pipeline sheets: {a.link}")
    if src.title == a.to:
        raise Guarded(2, f"job already in {a.to}")
    if a.to in STAGE_ORDER and STAGE_ORDER.get(src.title, 0) > STAGE_ORDER[a.to] and not a.allow_backward:
        raise Guarded(2, f"backward move {src.title} -> {a.to} refused; pass --allow-backward to reopen")
    if a.to == "Excluded" and src.title != "For Approval" and not a.allow_backward:
        raise Guarded(2, "only For Approval rows move to Excluded; close applied rows with --to Closed")
    hs = headers(src)
    row = {name: src.cell(r, col).value for name, col in hs.items()}
    dst = t.ws(a.to)
    hd = headers(dst)
    when = a.date or today()
    hist = str(row.get("Stage history") or "").strip("; ")
    hist = f"{hist}; {when} {a.to}" if hist else f"{when} {a.to}"
    if a.detail:
        hist += f" ({a.detail})"
    vals = {name: row.get(name) for name in hd if name != "#"}
    vals.update({"Stage history": hist, "Status date": when, "Latest status": a.detail or a.to,
                 "Pipeline stage": a.to})
    if a.to == "Applied":
        vals.update({"Date applied": when, "Follow-up date": add_days(when, a.follow_up_days),
                     "Channel": a.channel or row.get("Apply method"), "Fit": row.get("Fit score")})
    elif a.to == "Interviewing":
        vals.update({"Interview stage": a.interview_stage, "Next interview date": a.interview_date})
    elif a.to == "Closed":
        vals.update({"Closed date": when, "Outcome": a.outcome or "Closed", "Outcome detail": a.detail,
                     "Stage reached": src.title if src.title != "For Approval" else "Not applied"})
    elif a.to == "Excluded":
        notes = row.get("Decision notes")
        vals = {"Company": row.get("Company"), "Title": row.get("Title"), "Job link": row.get("Job link"),
                "Reason": (a.detail or "Declined by owner") + (f" - {notes}" if notes else ""),
                "Fit": row.get("Fit score"), "Pay": row.get("Pay (as posted)")}
    dr = last_row(dst, hd.get("Company", 1)) + 1
    copy_row_style(dst, dr - 1, dr)
    if "#" in hd:
        prev = dst.cell(dr - 1, hd["#"]).value
        vals["#"] = prev + 1 if isinstance(prev, int) else dr - 1
    put(dst, hd, dr, vals)
    src.delete_rows(r, 1)
    sm = sheet(t.wb, "Skill Map")
    if sm is not None:
        h = headers(sm)
        for i in range(2, sm.max_row + 1):
            if ct_key(get(sm, h, i, "Company"), get(sm, h, i, "Title")) == ct_key(row.get("Company"), row.get("Title")) \
                    and get(sm, h, i, "Bucket") == src.title:
                sm.cell(i, h["Bucket"]).value = a.to
    t.save(a.dry_run)
    return t.result({"moved": a.link, "from": src.title, "to": a.to, "row": dr}, a.dry_run)


AUTO_ACTION = re.compile(r"^(Close GAP|Advance|Review|Covered)\b")


def cmd_gaps(a) -> dict:
    """Classify Skill Map rows, rebuild Gap Summary, refresh readiness on open pipeline rows.

    rules JSON: {"cv_skills": [regex...], "rules": [{"match": regex, "ref": "ID", "status": "...",
                 "week": "...", "note": "..."}]}. First matching rule wins; no match => GAP.
    Owner-set Status values are kept unless --recompute.
    """
    cfg = json.loads(Path(a.rules).read_text(encoding="utf-8")) if a.rules else {}
    cv = [re.compile(x, re.I) for x in cfg.get("cv_skills", [])]
    rules = [(re.compile(x["match"], re.I), x) for x in cfg.get("rules", [])]
    t = Tracker(a.tracker, a.expect_mtime, a.backup_dir)
    sm = t.ws("Skill Map")
    h = headers(sm)
    per_job: dict = {}
    summary: dict = {}
    for r in range(2, sm.max_row + 1):
        skill = get(sm, h, r, "Skill")
        if skill in (None, ""):
            continue
        status, ref, week = get(sm, h, r, "Status"), get(sm, h, r, "Study plan ref"), get(sm, h, r, "Week/phase")
        if a.recompute or status in (None, ""):
            if any(rx.search(skill) for rx in cv):
                status, ref, week = CV, None, None
            else:
                hit = next((x for rx, x in rules if rx.search(skill)), None)
                status = hit.get("status", "Study plan - scheduled") if hit else "GAP - not in plan"
                ref, week = (hit.get("ref"), hit.get("week")) if hit else (None, None)
            put(sm, h, r, {"Status": status, "Study plan ref": ref, "Week/phase": week}, clear=True)
        bucket = get(sm, h, r, "Bucket")
        key = ct_key(get(sm, h, r, "Company"), get(sm, h, r, "Title"))
        per_job.setdefault(key, []).append((skill, str(status or ""), ref, week))
        if bucket in ("For Approval", "Applied", "Interviewing"):
            s = summary.setdefault(skill.strip().lower(), {"skill": skill, "jobs": set(), "status": status, "ref": ref})
            s["jobs"].add(key)
    gs = sheet(t.wb, "Gap Summary")
    if gs is not None:
        hg = headers(gs)
        gs.delete_rows(2, max(gs.max_row - 1, 0))
        rows = sorted(summary.values(), key=lambda s: (not str(s["status"]).startswith("GAP"), -len(s["jobs"]), s["skill"]))
        for i, s in enumerate(rows, start=2):
            gap = str(s["status"]).startswith("GAP")
            put(gs, hg, i, {"Skill": s["skill"], "# jobs requiring": len(s["jobs"]), "Status": s["status"],
                            "Study plan ref": s["ref"],
                            "Suggested action": "Add to study plan" if gap and len(s["jobs"]) >= a.gap_threshold
                            else ("Monitor" if gap else None)})
    updated = 0
    for name in ("For Approval", "Applied", "Interviewing"):
        ws = sheet(t.wb, name)
        if ws is None:
            continue
        hw = headers(ws)
        action_col = "Next study action" if "Next study action" in hw else "Next interview-prep action"
        for r in range(2, ws.max_row + 1):
            skills = per_job.get(ct_key(get(ws, hw, r, "Company"), get(ws, hw, r, "Title")))
            if not skills:
                continue
            scored = [x for x in skills if not x[1].lower().startswith("n/a")]
            if not scored:
                continue
            cov = [s for s, st, _, _ in scored if st == CV]
            gap = [s for s, st, _, _ in scored if st.startswith("GAP")]
            plan = [(s, ref, wk) for s, st, ref, wk in scored if st != CV and not st.startswith("GAP")
                    and (ref or st.lower().startswith(("study plan", "plan")))]
            ready = round((len(cov) + 0.5 * len(plan)) / len(scored), 2)
            fmt = [f"{s} ({ref}{', ' + wk if wk else ''})" if ref else s for s, ref, wk in plan]
            put(ws, hw, r, {"Readiness %": ready, "Covered": "; ".join(cov) or None, "In study plan": "; ".join(fmt) or None,
                            "Missing (not in plan)": "; ".join(gap) or None})
            cur = get(ws, hw, r, action_col)
            if action_col in hw and (cur in (None, "") or AUTO_ACTION.match(str(cur))):
                nxt = f"Close GAP: {gap[0]} - not in study plan" if gap else (f"Advance: {fmt[0]}" if fmt else "Covered - review stories")
                ws.cell(r, hw[action_col]).value = nxt
            if "Readiness %" in hw:
                ws.cell(r, hw["Readiness %"]).number_format = "0%"
            updated += 1
    t.save(a.dry_run)
    gaps = [{"skill": s["skill"], "jobs": len(s["jobs"])} for s in summary.values()
            if str(s["status"]).startswith("GAP") and len(s["jobs"]) >= a.gap_threshold]
    return t.result({"job_rows_updated": updated, "repeated_gaps": sorted(gaps, key=lambda g: -g["jobs"])}, a.dry_run)


SIDEGIG_FIELDS = {"platform": "Platform", "link": "Link", "status": "My status", "applied": "Roles/projects applied",
                  "pay": "Pay (USD/hr)", "meets_floor": "Meets my floor", "eligible": "Eligible from my country",
                  "payment": "Payment method", "next_action": "Next action", "next_action_date": "Next action date",
                  "last_activity": "Last activity", "mail_link": "Mail link", "notes": "Notes/sources"}


def cmd_sidegig(a) -> dict:
    """Upsert Side Gigs rows by Platform from a JSON list of {platform, status, next_action, ...}."""
    data = json.loads(Path(a.data).read_text(encoding="utf-8"))
    items = data["platforms"] if isinstance(data, dict) else data
    t = Tracker(a.tracker, a.expect_mtime, a.backup_dir)
    ws = t.ws("Side Gigs")
    h = headers(ws)
    if "Platform" not in h:
        raise Guarded(2, "Side Gigs sheet has no 'Platform' column")
    rows = {str(get(ws, h, r, "Platform")).strip().lower(): r for r in range(2, ws.max_row + 1)
            if get(ws, h, r, "Platform") not in (None, "")}
    updated, added = [], []
    for it in items:
        name = str(it.get("platform") or "").strip()
        if not name:
            raise Guarded(2, f"side gig entry missing 'platform': {it}")
        vals = {SIDEGIG_FIELDS[k]: v for k, v in it.items() if k in SIDEGIG_FIELDS}
        r = rows.get(name.lower())
        if r is None:
            r = last_row(ws, h["Platform"]) + 1
            copy_row_style(ws, r - 1, r)
            if "#" in h:
                prev = ws.cell(r - 1, h["#"]).value
                vals["#"] = prev + 1 if isinstance(prev, int) else r - 1
            rows[name.lower()] = r
            added.append(name)
        else:
            updated.append(name)
        put(ws, h, r, vals)
    t.save(a.dry_run)
    return t.result({"updated": updated, "added": added,
                     "ignored_keys": sorted({k for it in items for k in it if k not in SIDEGIG_FIELDS})}, a.dry_run)


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)

    def common(p, write=False):
        p.add_argument("--tracker", required=True)
        p.add_argument("--id-pattern", action="append", help="regex with one group extracting a job id from a link")
        if write:
            p.add_argument("--expect-mtime", type=float, help="abort (exit 3) unless the file still has this mtime")
            p.add_argument("--backup-dir", help="copy the workbook here before writing")
            p.add_argument("--dry-run", action="store_true")
        return p

    p = sub.add_parser("init", help="create an empty tracker with the standard sheets")
    p.add_argument("--tracker", required=True)
    p.add_argument("--force", action="store_true")
    common(sub.add_parser("known", help="dump known job ids and company|title keys (dedup list)")).add_argument("--out")
    common(sub.add_parser("approvals", help="list For Approval rows by owner mark"))
    p = common(sub.add_parser("append", help="append screened jobs to For Approval + Skill Map"), True)
    p.add_argument("--jobs", required=True)
    p.add_argument("--note", default=f"New {today()} scan")
    p.add_argument("--cv-skill", action="append", help="regex; matching skills get Status 'Experience (CV)'")
    p = common(sub.add_parser("exclude", help="append screened-out jobs to Excluded"), True)
    p.add_argument("--jobs", required=True)
    p = common(sub.add_parser("move", help="move a job between pipeline stages"), True)
    p.add_argument("--link", required=True)
    p.add_argument("--to", required=True, choices=["Applied", "Interviewing", "Closed", "Excluded"])
    p.add_argument("--date")
    p.add_argument("--detail")
    p.add_argument("--channel")
    p.add_argument("--outcome")
    p.add_argument("--interview-stage")
    p.add_argument("--interview-date")
    p.add_argument("--follow-up-days", type=int, default=7)
    p.add_argument("--allow-backward", action="store_true", help="permit reopening a later-stage row")
    p = common(sub.add_parser("gaps", help="classify skills, rebuild Gap Summary, refresh readiness"), True)
    p.add_argument("--rules")
    p.add_argument("--recompute", action="store_true")
    p.add_argument("--gap-threshold", type=int, default=2)
    p = common(sub.add_parser("sidegig", help="upsert Side Gigs rows by platform"), True)
    p.add_argument("--data", required=True)
    a = ap.parse_args(argv)
    fn = {"init": cmd_init, "known": cmd_known, "approvals": cmd_approvals, "append": cmd_append,
          "exclude": cmd_exclude, "move": cmd_move, "gaps": cmd_gaps, "sidegig": cmd_sidegig}[a.cmd]
    try:
        out = fn(a)
    except Guarded as e:
        sys.stderr.write(f"tracker_ops: {e}\n")
        return e.code
    sys.stdout.write(json.dumps(out, ensure_ascii=False, indent=1, default=str) + "\n")
    return 0


if __name__ == "__main__":
    sys.exit(main())
